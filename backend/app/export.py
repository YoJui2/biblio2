"""Generación del libro de Excel exportable del catálogo.

El archivo se arma con tres hojas pensadas para leerse sin ayuda:

* ``Catálogo``: una fila por registro bibliográfico, con el formato visual
  del panel (verde, filas alternadas, filtros y paneles congelados).
* ``Resumen``: totales, rango de inventario y rankings de autores y títulos.
* ``Ejemplares``: una fila por cada número de inventario, útil para buscar
  un ejemplar puntual.
"""

from __future__ import annotations

from collections import Counter
from datetime import datetime
from io import BytesIO
from typing import Any, Iterable

from openpyxl import Workbook
from openpyxl.formatting.rule import DataBarRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.worksheet import Worksheet

# Paleta alineada con styles.css del panel.
GREEN_DARK = "315F45"
GREEN = "3C7453"
GREEN_SOFT = "E8F0E9"
GREEN_PALE = "F3F7F2"
ZEBRA = "F8FAF7"
INK = "20231F"
MUTED = "6E746C"
LINE = "E2E8DF"
LINE_SOFT = "EDF1EA"
WHITE = "FFFFFF"

TITLE_FONT = "Calibri"
BODY_FONT = "Calibri"

THIN = Side(style="thin", color=LINE)
HAIR = Side(style="hair", color=LINE_SOFT)
CELL_BORDER = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)
TOTAL_BORDER = Border(top=Side(style="medium", color=GREEN), bottom=THIN, left=HAIR, right=HAIR)

CATALOG_SHEET = "Catálogo"
SUMMARY_SHEET = "Resumen"
COPIES_SHEET = "Ejemplares"
CATALOG_COLUMNS = 6


def format_inventory_number(number: int) -> str:
    return f"INV-{number:06d}"


def inventory_text(numbers: Iterable[int]) -> str:
    """Compacta números consecutivos en rangos, igual que el panel."""
    ordered = sorted(set(numbers))
    if not ordered:
        return "Sin número"
    ranges: list[str] = []
    start = previous = ordered[0]
    for number in ordered[1:]:
        if number == previous:
            continue
        if number == previous + 1:
            previous = number
            continue
        ranges.append(_range_text(start, previous))
        start = previous = number
    ranges.append(_range_text(start, previous))
    return ", ".join(ranges)


def _range_text(start: int, end: int) -> str:
    if start == end:
        return format_inventory_number(start)
    return f"{format_inventory_number(start)}–{format_inventory_number(end)}"


def _searchable_text(record: dict[str, Any]) -> str:
    parts = [
        record.get("title", ""),
        record.get("author", ""),
        record.get("imprint", ""),
        record.get("binding", ""),
        record.get("inventory_source", ""),
        *(format_inventory_number(number) for number in record.get("inventory_numbers", [])),
    ]
    return " ".join(str(part) for part in parts).casefold()


def filter_records(
    records: Iterable[dict[str, Any]], query: str
) -> list[dict[str, Any]]:
    needle = (query or "").strip().casefold()
    if not needle:
        return list(records)
    return [record for record in records if needle in _searchable_text(record)]


def build_catalog_workbook(
    records: Iterable[dict[str, Any]],
    *,
    query: str = "",
    generated_at: datetime | None = None,
) -> bytes:
    """Devuelve el contenido binario de un ``.xlsx`` listo para descargar."""
    selected = filter_records(records, query)
    generated_at = generated_at or datetime.now()

    workbook = Workbook()
    catalog = workbook.active
    catalog.title = CATALOG_SHEET
    _build_catalog_sheet(catalog, selected, query, generated_at)
    _build_summary_sheet(workbook.create_sheet(SUMMARY_SHEET), selected, query, generated_at)
    _build_copies_sheet(workbook.create_sheet(COPIES_SHEET), selected, generated_at)

    workbook.properties.title = "Inventario de la biblioteca"
    workbook.properties.subject = "Catálogo bibliográfico y ejemplares"
    workbook.properties.creator = "Inventario de Biblioteca"
    workbook.properties.created = generated_at

    buffer = BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def _build_catalog_sheet(
    sheet: Worksheet,
    records: list[dict[str, Any]],
    query: str,
    generated_at: datetime,
) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = GREEN
    _title_band(
        sheet,
        "Inventario de la biblioteca",
        _subtitle(len(records), query, generated_at),
        total_columns=CATALOG_COLUMNS,
    )

    copies = sum(len(record["inventory_numbers"]) for record in records)
    numbers = [number for record in records for number in record["inventory_numbers"]]
    info_row = 4
    info_values = [
        (1, 2, f"{len(records):,} registros bibliográficos".replace(",", ".")),
        (3, 4, f"{copies:,} ejemplares numerados".replace(",", ".")),
        (
            5,
            6,
            f"Rango {format_inventory_number(min(numbers))} – {format_inventory_number(max(numbers))}"
            if numbers
            else "Sin ejemplares registrados",
        ),
    ]
    for first, last, text in info_values:
        sheet.merge_cells(start_row=info_row, start_column=first, end_row=info_row, end_column=last)
        cell = sheet.cell(row=info_row, column=first, value=text)
        fill = PatternFill("solid", fgColor=GREEN_SOFT)
        cell.fill = fill
        cell.font = Font(name=BODY_FONT, size=10, bold=True, color=GREEN_DARK)
        cell.alignment = Alignment(vertical="center", horizontal="center")
        _paint_merged(sheet, info_row, first, last, fill)
        _outline_merged(sheet, info_row, first, last, Side(style="thin", color="C9DCC9"))
    sheet.row_dimensions[info_row].height = 26
    sheet.row_dimensions[5].height = 8

    headers = [
        "N° de inventario",
        "Título",
        "Autor",
        "Pie de imprenta",
        "Ejemplares",
        "Encuadernación / Observaciones",
    ]
    header_row = 6
    _header_row(sheet, header_row, headers)
    widths = [30, 46, 29, 38, 11, 32]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    first_data_row = header_row + 1
    for offset, record in enumerate(records):
        row = first_data_row + offset
        zebra = offset % 2 == 1
        values = [
            inventory_text(record.get("inventory_numbers", [])),
            record.get("title") or "—",
            record.get("author") or "—",
            record.get("imprint") or "—",
            len(record.get("inventory_numbers", [])),
            record.get("binding") or "—",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = CELL_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=column in (1, 2, 4, 6))
            cell.font = Font(name=BODY_FONT, size=10, color=INK)
            if zebra:
                cell.fill = PatternFill("solid", fgColor=ZEBRA)
        sheet.cell(row=row, column=1).font = Font(name="Consolas", size=9, color=GREEN_DARK)
        sheet.cell(row=row, column=2).font = Font(name=BODY_FONT, size=10, bold=True, color=INK)
        sheet.cell(row=row, column=3).font = Font(name=BODY_FONT, size=10, color=MUTED)
        sheet.cell(row=row, column=4).font = Font(name=BODY_FONT, size=9, color=MUTED)
        sheet.cell(row=row, column=5).alignment = Alignment(vertical="center", horizontal="center")
        sheet.cell(row=row, column=5).font = Font(name=BODY_FONT, size=10, bold=True, color=GREEN_DARK)
        sheet.cell(row=row, column=6).font = Font(name=BODY_FONT, size=9, color=MUTED)
        sheet.row_dimensions[row].height = 34

    last_data_row = first_data_row + len(records) - 1 if records else header_row

    total_row = last_data_row + 1
    sheet.merge_cells(start_row=total_row, start_column=1, end_row=total_row, end_column=3)
    label = sheet.cell(row=total_row, column=1, value="Total del catálogo exportado")
    sheet.cell(row=total_row, column=5, value=copies)
    for column in range(1, CATALOG_COLUMNS + 1):
        cell = sheet.cell(row=total_row, column=column)
        cell.fill = PatternFill("solid", fgColor=GREEN_PALE)
        cell.border = TOTAL_BORDER
    label.font = Font(name=BODY_FONT, size=10, bold=True, color=GREEN_DARK)
    label.alignment = Alignment(vertical="center", horizontal="right")
    total_cell = sheet.cell(row=total_row, column=5)
    total_cell.font = Font(name=BODY_FONT, size=11, bold=True, color=GREEN_DARK)
    total_cell.alignment = Alignment(vertical="center", horizontal="center")
    sheet.row_dimensions[total_row].height = 28

    if records:
        sheet.conditional_formatting.add(
            f"E{first_data_row}:E{last_data_row}",
            DataBarRule(
                start_type="num",
                start_value=0,
                end_type="max",
                color="9CC09C",
                showValue=True,
            ),
        )
        sheet.auto_filter.ref = f"A{header_row}:F{last_data_row}"
    sheet.freeze_panes = f"A{first_data_row}"

    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.paperSize = sheet.PAPERSIZE_A4
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = f"1:{header_row}"
    sheet.print_options.horizontalCentered = True


def _build_summary_sheet(
    sheet: Worksheet,
    records: list[dict[str, Any]],
    query: str,
    generated_at: datetime,
) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "80A26F"
    sheet.column_dimensions["A"].width = 38
    sheet.column_dimensions["B"].width = 24
    sheet.column_dimensions["C"].width = 18
    sheet.column_dimensions["D"].width = 16

    _title_band(
        sheet,
        "Resumen del inventario",
        "Indicadores, rango y concentración del catálogo exportado.",
        total_columns=4,
    )

    copies = [number for record in records for number in record["inventory_numbers"]]
    average = (sum(len(record["inventory_numbers"]) for record in records) / len(records)) if records else 0
    _section_title(sheet, 4, "Indicadores generales", 4)

    indicators = [
        ("Registros bibliográficos", len(records)),
        ("Ejemplares numerados", len(copies)),
        ("Promedio de ejemplares por registro", f"{average:.1f}"),
        ("Registros sin número de inventario", sum(1 for record in records if not record["inventory_numbers"])),
        ("Primer número de inventario", format_inventory_number(min(copies)) if copies else "—"),
        ("Último número de inventario", format_inventory_number(max(copies)) if copies else "—"),
        ("Filtro de búsqueda aplicado", query.strip() or "Sin filtro (catálogo completo)"),
        ("Generado", generated_at.strftime("%d/%m/%Y %H:%M")),
    ]
    row = 5
    for label, value in indicators:
        text_value = not isinstance(value, int) and not (
            isinstance(value, str) and value.replace(".", "", 1).isdigit()
        )
        _metric_row(sheet, row, label, value, text_value=text_value)
        row += 1

    row += 1
    _section_title(sheet, row, "Autores con más registros", 4)
    row += 1
    _ranking_header(sheet, row, ["#", "Autor", "Registros", "Ejemplares"])
    registers_by_author: Counter[str] = Counter()
    copies_by_author: Counter[str] = Counter()
    for record in records:
        author = (record.get("author") or "").strip() or "Sin autor"
        registers_by_author[author] += 1
        copies_by_author[author] += len(record["inventory_numbers"])
    row += 1
    for position, (author, count) in enumerate(registers_by_author.most_common(10), start=1):
        _ranking_row(
            sheet,
            row,
            [position, author, count, copies_by_author[author]],
            zebra=position % 2 == 0,
        )
        row += 1
    if not records:
        _empty_note(sheet, row, "No hay registros para resumir.", 4)
        row += 1

    row += 1
    _section_title(sheet, row, "Títulos con más ejemplares", 4)
    row += 1
    _ranking_header(sheet, row, ["#", "Título", "Ejemplares", "Autor"])
    row += 1
    top_titles = sorted(
        (record for record in records if record["inventory_numbers"]),
        key=lambda record: (len(record["inventory_numbers"]), record.get("title") or ""),
        reverse=True,
    )[:10]
    for position, record in enumerate(top_titles, start=1):
        _ranking_row(
            sheet,
            row,
            [
                position,
                record.get("title") or "Sin título",
                len(record["inventory_numbers"]),
                record.get("author") or "Sin autor",
            ],
            zebra=position % 2 == 0,
        )
        row += 1
    if not top_titles:
        _empty_note(sheet, row, "No hay títulos con ejemplares para resumir.", 4)

    sheet.freeze_panes = "A4"
    sheet.page_setup.orientation = "portrait"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True


def _build_copies_sheet(
    sheet: Worksheet,
    records: list[dict[str, Any]],
    generated_at: datetime,
) -> None:
    sheet.sheet_view.showGridLines = False
    sheet.sheet_properties.tabColor = "B9CDB4"
    headers = ["N° de inventario", "Título", "Autor", "Pie de imprenta", "Encuadernación / Observaciones"]
    _title_band(
        sheet,
        "Ejemplares numerados",
        f"Una fila por número de inventario · Generado el {generated_at.strftime('%d/%m/%Y %H:%M')}",
        total_columns=len(headers),
    )
    _header_row(sheet, 4, headers)
    widths = [20, 44, 28, 36, 30]
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = width

    copies = sorted(
        (
            (number, record)
            for record in records
            for number in record["inventory_numbers"]
        ),
        key=lambda item: item[0],
    )
    first_data_row = 5
    for offset, (number, record) in enumerate(copies):
        row = first_data_row + offset
        zebra = offset % 2 == 1
        values = [
            format_inventory_number(number),
            record.get("title") or "—",
            record.get("author") or "—",
            record.get("imprint") or "—",
            record.get("binding") or "—",
        ]
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=row, column=column, value=value)
            cell.border = CELL_BORDER
            cell.alignment = Alignment(vertical="center", wrap_text=column in (2, 4, 5))
            cell.font = Font(name=BODY_FONT, size=10, color=INK)
            if zebra:
                cell.fill = PatternFill("solid", fgColor=ZEBRA)
        sheet.cell(row=row, column=1).font = Font(name="Consolas", size=9, color=GREEN_DARK)
        sheet.cell(row=row, column=2).font = Font(name=BODY_FONT, size=10, bold=True, color=INK)
        sheet.cell(row=row, column=3).font = Font(name=BODY_FONT, size=10, color=MUTED)
        sheet.cell(row=row, column=4).font = Font(name=BODY_FONT, size=9, color=MUTED)
        sheet.cell(row=row, column=5).font = Font(name=BODY_FONT, size=9, color=MUTED)
        sheet.row_dimensions[row].height = 30

    if copies:
        sheet.auto_filter.ref = f"A4:E{first_data_row + len(copies) - 1}"
    sheet.freeze_panes = f"A{first_data_row}"
    sheet.page_setup.orientation = "landscape"
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.print_title_rows = "1:4"


def _title_band(sheet: Worksheet, title: str, subtitle: str, *, total_columns: int) -> None:
    sheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=total_columns)
    title_cell = sheet.cell(row=1, column=1, value=title)
    title_cell.font = Font(name=TITLE_FONT, size=18, bold=True, color=WHITE)
    title_cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    title_fill = PatternFill("solid", fgColor=GREEN_DARK)
    _paint_merged(sheet, 1, 1, total_columns, title_fill)
    sheet.row_dimensions[1].height = 42

    sheet.merge_cells(start_row=2, start_column=1, end_row=2, end_column=total_columns)
    subtitle_cell = sheet.cell(row=2, column=1, value=subtitle)
    subtitle_cell.font = Font(name=BODY_FONT, size=10, color=GREEN_SOFT)
    subtitle_cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    subtitle_fill = PatternFill("solid", fgColor=GREEN)
    _paint_merged(sheet, 2, 1, total_columns, subtitle_fill)
    sheet.row_dimensions[2].height = 24

    sheet.row_dimensions[3].height = 8


def _subtitle(count: int, query: str, generated_at: datetime) -> str:
    suffix = f" · Filtro: «{query.strip()}»" if query.strip() else " · Catálogo completo"
    return (
        f"{count} registro(s){suffix} · Actualizado el "
        f"{generated_at.strftime('%d/%m/%Y a las %H:%M')}"
    )


def _header_row(sheet: Worksheet, row: int, headers: list[str]) -> None:
    fill = PatternFill("solid", fgColor=GREEN)
    for column, text in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=text)
        cell.fill = fill
        cell.font = Font(name=BODY_FONT, size=10, bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center", horizontal="left", wrap_text=True, indent=1)
        cell.border = Border(left=Side(style="thin", color=GREEN), right=Side(style="thin", color=GREEN))
    sheet.row_dimensions[row].height = 28


def _section_title(sheet: Worksheet, row: int, text: str, total_columns: int) -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_columns)
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = Font(name=BODY_FONT, size=12, bold=True, color=GREEN_DARK)
    cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)
    _paint_merged(sheet, row, 1, total_columns, PatternFill("solid", fgColor=GREEN_SOFT))
    sheet.row_dimensions[row].height = 26


def _metric_row(sheet: Worksheet, row: int, label: str, value: Any, *, text_value: bool) -> None:
    label_cell = sheet.cell(row=row, column=1, value=label)
    label_cell.font = Font(name=BODY_FONT, size=10, color=MUTED)
    label_cell.alignment = Alignment(vertical="center", indent=1)
    label_cell.border = Border(bottom=HAIR)

    value_cell = sheet.cell(row=row, column=2, value=value)
    value_cell.font = Font(name=BODY_FONT, size=10 if text_value else 12, bold=True, color=INK)
    value_cell.alignment = Alignment(vertical="center", horizontal="left" if text_value else "right", indent=0 if text_value else 1)
    value_cell.border = Border(bottom=HAIR)
    sheet.row_dimensions[row].height = 22


def _ranking_header(sheet: Worksheet, row: int, headers: list[str]) -> None:
    for column, text in enumerate(headers, start=1):
        cell = sheet.cell(row=row, column=column, value=text)
        cell.fill = PatternFill("solid", fgColor=GREEN_DARK)
        cell.font = Font(name=BODY_FONT, size=9, bold=True, color=WHITE)
        cell.alignment = Alignment(vertical="center", horizontal="center" if column != 2 else "left", indent=0 if column != 2 else 1)
        cell.border = Border(bottom=THIN)
    sheet.row_dimensions[row].height = 22


def _ranking_row(sheet: Worksheet, row: int, values: list[Any], *, zebra: bool) -> None:
    fill = PatternFill("solid", fgColor=ZEBRA) if zebra else None
    for column, value in enumerate(values, start=1):
        cell = sheet.cell(row=row, column=column, value=value)
        cell.font = Font(name=BODY_FONT, size=10, color=INK, bold=column == 2)
        cell.alignment = Alignment(vertical="center", horizontal="center" if column in (1, 3, 4) else "left", indent=0 if column in (1, 3, 4) else 1)
        cell.border = Border(bottom=HAIR)
        if fill is not None:
            cell.fill = fill
    sheet.row_dimensions[row].height = 21


def _empty_note(sheet: Worksheet, row: int, text: str, total_columns: int) -> None:
    sheet.merge_cells(start_row=row, start_column=1, end_row=row, end_column=total_columns)
    cell = sheet.cell(row=row, column=1, value=text)
    cell.font = Font(name=BODY_FONT, size=10, italic=True, color=MUTED)
    cell.alignment = Alignment(vertical="center", horizontal="left", indent=1)


def _paint_merged(sheet: Worksheet, row: int, first: int, last: int, fill: PatternFill) -> None:
    for column in range(first, last + 1):
        sheet.cell(row=row, column=column).fill = fill


def _outline_merged(sheet: Worksheet, row: int, first: int, last: int, side: Side) -> None:
    for column in range(first, last + 1):
        cell = sheet.cell(row=row, column=column)
        left = side if column == first else Side()
        right = side if column == last else Side()
        cell.border = Border(left=left, right=right, top=side, bottom=side)
