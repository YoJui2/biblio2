from io import BytesIO

from fastapi.testclient import TestClient
from openpyxl import load_workbook
import pytest

from app.main import create_app


@pytest.fixture
def client(tmp_path):
    database_url = f"sqlite+pysqlite:///{tmp_path / 'test-export.db'}"
    application = create_app(database_url, create_schema=True)
    with TestClient(application) as test_client:
        yield test_client


def _seed(client: TestClient) -> None:
    client.post(
        "/api/books",
        json={
            "title": "El Aleph",
            "author": "Jorge Luis Borges",
            "imprint": "Buenos Aires: Emecé, 1949",
            "binding": "Cartoné",
            "inventory_numbers": [62, 63, 64],
            "inventory_source": "62-64",
        },
    )
    client.post(
        "/api/books",
        json={
            "title": "Rayuela",
            "author": "Julio Cortázar",
            "imprint": "Buenos Aires: Sudamericana, 1963",
            "binding": "",
            "inventory_numbers": [100],
            "inventory_source": "100",
        },
    )


def test_export_xlsx_has_formatted_catalog_summary_and_copies(client: TestClient) -> None:
    _seed(client)

    response = client.get("/api/export.xlsx")

    assert response.status_code == 200
    assert response.headers["content-type"].startswith(
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    )
    assert "inventario-biblioteca" in response.headers["content-disposition"]

    workbook = load_workbook(BytesIO(response.content))
    assert workbook.sheetnames == ["Catálogo", "Resumen", "Ejemplares"]

    catalog = workbook["Catálogo"]
    assert catalog["A1"].value == "Inventario de la biblioteca"
    assert catalog["A6"].value == "N° de inventario"
    assert catalog["A7"].value == "INV-000062–INV-000064"
    assert catalog["B7"].value == "El Aleph"
    assert catalog["E7"].value == 3
    assert catalog["B8"].value == "Rayuela"
    assert catalog.freeze_panes == "A7"
    assert catalog.auto_filter.ref == "A6:F8"
    assert catalog["A9"].value == "Total del catálogo exportado"
    assert catalog["E9"].value == 4
    assert catalog.sheet_view.showGridLines is False

    summary = workbook["Resumen"]
    values = [cell.value for row in summary.iter_rows() for cell in row if cell.value]
    assert "Registros bibliográficos" in values
    assert "Jorge Luis Borges" in values
    assert "Rayuela" in values

    copies = workbook["Ejemplares"]
    assert copies["A4"].value == "N° de inventario"
    assert [copies.cell(row=row, column=1).value for row in range(5, 9)] == [
        "INV-000062",
        "INV-000063",
        "INV-000064",
        "INV-000100",
    ]
    assert copies.freeze_panes == "A5"


def test_export_xlsx_respects_the_search_filter(client: TestClient) -> None:
    _seed(client)

    response = client.get("/api/export.xlsx", params={"q": "cortázar"})

    assert response.status_code == 200
    assert "filtrado" in response.headers["content-disposition"]
    catalog = load_workbook(BytesIO(response.content))["Catálogo"]
    assert catalog["B7"].value == "Rayuela"
    assert catalog["B8"].value is None
    assert "Filtro: «cortázar»" in catalog["A2"].value
    assert catalog["E8"].value == 1


def test_export_xlsx_with_empty_catalog_still_builds(client: TestClient) -> None:
    response = client.get("/api/export.xlsx")

    assert response.status_code == 200
    workbook = load_workbook(BytesIO(response.content))
    assert workbook["Catálogo"]["A7"].value == "Total del catálogo exportado"
    assert workbook["Catálogo"]["E7"].value == 0
    assert workbook["Ejemplares"].max_row == 4
